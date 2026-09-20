# -*- coding: utf-8 -*-
"""
Table 1 — Patient Characteristics (FULL feature set)
═══════════════════════════════════════════════════════

对所有特征同时算两组对比:
  - DEV (n=278) vs VAL (n=70):  验证 80/20 划分的同质性 (p 应不显著)
  - SSI (n≈55) vs non-SSI (n≈293):  单变量风险因子识别

输出:
  Table1_patient_characteristics.xlsx   (主输出,易于编辑)
  Table1_patient_characteristics.csv    (备用)

数据来源:
  开腹肝胆道结石手术SSI清洗.csv
  splits/split_folds_full.csv
  特征类别/feature_types_summary.csv

使用方法:
  把此脚本放在项目根目录 (或 tables/ 子目录,会自动找上一级)
  python make_table1.py

依赖:
  pandas, numpy, scipy, openpyxl
"""

import sys
from pathlib import Path

# ── 自动定位项目根目录 ─────────────────────────────────
# 优先使用脚本所在目录,如果是子目录 (如 tables/) 则向上一级
SCRIPT_DIR = Path(__file__).resolve().parent
if (SCRIPT_DIR / "开腹肝胆道结石手术SSI清洗.csv").exists():
    ROOT = SCRIPT_DIR
elif (SCRIPT_DIR.parent / "开腹肝胆道结石手术SSI清洗.csv").exists():
    ROOT = SCRIPT_DIR.parent
else:
    raise FileNotFoundError(
        "找不到 开腹肝胆道结石手术SSI清洗.csv\n"
        "把此脚本放在项目根目录,或在 tables/ 子目录下也可"
    )

import pandas as pd
import numpy as np
from scipy import stats

# ── 路径 ──────────────────────────────────────────────
DATA_CSV          = ROOT / "开腹肝胆道结石手术SSI清洗.csv"
SPLIT_CSV         = ROOT / "splits/split_folds_full.csv"
FEATURE_TYPES_CSV = ROOT / "特征类别/feature_types_summary.csv"

OUT_DIR  = ROOT / "Table"
OUT_DIR.mkdir(exist_ok=True)
OUT_XLSX = OUT_DIR / "Table1_patient_characteristics.xlsx"
OUT_CSV  = OUT_DIR / "Table1_patient_characteristics.csv"


# ══════════════════════════════════════════════════════
# 1. Load and merge
# ══════════════════════════════════════════════════════
print(f"[INFO] ROOT = {ROOT}")

df = pd.read_csv(DATA_CSV)
sp = pd.read_csv(SPLIT_CSV)
print(f"[INFO] Loaded raw data: {df.shape}, split table: {sp.shape}")

df = df.merge(
    sp[["ID", "split", "Infection"]],
    on="ID", how="inner", suffixes=("", "_split"),
)
y_col = "Infection_split" if "Infection_split" in df.columns else "Infection"
df[y_col] = pd.to_numeric(df[y_col], errors="raise").astype(int)
df["split"] = df["split"].astype(str).str.upper()

# ── 加载特征类型 ──
if FEATURE_TYPES_CSV.exists():
    ft = pd.read_csv(FEATURE_TYPES_CSV)
    cont_cols = ft[ft["type"].str.lower() == "continuous"]["feature"].tolist()
    bin_cols  = ft[ft["type"].str.lower() == "binary"]["feature"].tolist()
    cat_cols  = ft[ft["type"].str.lower().isin(["categorical", "category"])]["feature"].tolist()
else:
    print("[WARN] feature_types_summary.csv 未找到,使用启发式判断")
    cont_cols, bin_cols, cat_cols = [], [], []
    for c in df.columns:
        if c in ("ID", "split", "fold_id", "Infection", y_col):
            continue
        u = df[c].dropna().unique()
        if len(u) == 2 and set(u).issubset({0, 1, 0.0, 1.0}):
            bin_cols.append(c)
        elif df[c].dtype in (np.float64, np.int64) and len(u) > 10:
            cont_cols.append(c)
        else:
            cat_cols.append(c)

# 排除非特征列
EXCLUDE = {"ID", "split", "fold_id", "Infection", y_col, "Infection_split"}
cont_cols = [c for c in cont_cols if c in df.columns and c not in EXCLUDE]
bin_cols  = [c for c in bin_cols  if c in df.columns and c not in EXCLUDE]
cat_cols  = [c for c in cat_cols  if c in df.columns and c not in EXCLUDE]

print(f"[INFO] Continuous: {len(cont_cols)} | Binary: {len(bin_cols)} | Categorical: {len(cat_cols)}")

# ══════════════════════════════════════════════════════
# 2. Helpers
# ══════════════════════════════════════════════════════
def fmt_cont(values):
    """median (IQR) for continuous variables."""
    v = pd.Series(values).dropna()
    if len(v) == 0:
        return "—"
    med = v.median()
    q1, q3 = v.quantile([0.25, 0.75])
    return f"{med:.2f} ({q1:.2f}–{q3:.2f})"


def fmt_np(positives, total):
    """n (%) — total excludes missings of that subgroup."""
    if total == 0:
        return "—"
    pct = positives / total * 100
    return f"{int(positives)} ({pct:.1f})"


def test_cont(g1, g2):
    """Mann-Whitney U, two-sided."""
    g1 = pd.Series(g1).dropna()
    g2 = pd.Series(g2).dropna()
    if len(g1) < 2 or len(g2) < 2:
        return np.nan
    # 全相同时 Mann-Whitney 退化
    if g1.nunique() == 1 and g2.nunique() == 1 and g1.iloc[0] == g2.iloc[0]:
        return 1.0
    try:
        _, p = stats.mannwhitneyu(g1, g2, alternative="two-sided")
        return p
    except Exception as e:
        return np.nan


def test_cat(g1_vals, g2_vals):
    """Chi-square,2x2 时若期望频数 < 5 自动切 Fisher."""
    g1 = pd.Series(g1_vals).dropna()
    g2 = pd.Series(g2_vals).dropna()
    levels = sorted(set(g1.unique()) | set(g2.unique()))
    if len(levels) < 2:
        return np.nan
    ct = np.array([
        [int((g1 == lv).sum()) for lv in levels],
        [int((g2 == lv).sum()) for lv in levels],
    ])
    if ct.sum() == 0:
        return np.nan
    # 全部集中在一列也无法检验
    col_sums = ct.sum(axis=0)
    if (col_sums == 0).any():
        return np.nan
    try:
        chi2, p, dof, expected = stats.chi2_contingency(ct)
        if ct.shape == (2, 2) and (expected < 5).any():
            _, p = stats.fisher_exact(ct)
        return p
    except Exception:
        return np.nan


def fmt_p(p):
    if pd.isna(p):
        return "—"
    if p < 0.001:
        return "<0.001"
    return f"{p:.3f}"


# ══════════════════════════════════════════════════════
# 3. Build rows
# ══════════════════════════════════════════════════════
total_n   = len(df)
dev_n     = (df["split"] == "DEV").sum()
val_n     = (df["split"] == "VAL").sum()
ssi_n     = (df[y_col] == 1).sum()
non_ssi_n = (df[y_col] == 0).sum()

print(f"[INFO] Total {total_n} | DEV {dev_n} | VAL {val_n} | SSI {ssi_n} | non-SSI {non_ssi_n}")

mask_dev = df["split"] == "DEV"
mask_val = df["split"] == "VAL"
mask_ssi = df[y_col] == 1
mask_non = df[y_col] == 0

rows = []

# Top-level n header row
rows.append({
    "Variable": "n",
    "Type": "",
    "All (n=%d)" % total_n: total_n,
    "DEV (n=%d)" % dev_n: dev_n,
    "VAL (n=%d)" % val_n: val_n,
    "p (DEV vs VAL)": "",
    "Non-SSI (n=%d)" % non_ssi_n: non_ssi_n,
    "SSI (n=%d)" % ssi_n: ssi_n,
    "p (SSI vs non-SSI)": "",
    "Missing n": "",
    "Missing %": "",
})


def add_row(feat, ftype, *,
            all_str, dev_str, val_str, p_split,
            non_str, ssi_str, p_ssi,
            missing_n):
    rows.append({
        "Variable": feat,
        "Type": ftype,
        "All (n=%d)" % total_n: all_str,
        "DEV (n=%d)" % dev_n: dev_str,
        "VAL (n=%d)" % val_n: val_str,
        "p (DEV vs VAL)": p_split,
        "Non-SSI (n=%d)" % non_ssi_n: non_str,
        "SSI (n=%d)" % ssi_n: ssi_str,
        "p (SSI vs non-SSI)": p_ssi,
        "Missing n": missing_n,
        "Missing %": f"{missing_n / total_n * 100:.2f}%" if missing_n is not None else "",
    })


# ── 3a. Continuous ──
for feat in cont_cols:
    s = df[feat]
    miss = int(s.isna().sum())
    add_row(
        feat, "continuous",
        all_str=fmt_cont(s),
        dev_str=fmt_cont(df.loc[mask_dev, feat]),
        val_str=fmt_cont(df.loc[mask_val, feat]),
        p_split=fmt_p(test_cont(df.loc[mask_dev, feat], df.loc[mask_val, feat])),
        non_str=fmt_cont(df.loc[mask_non, feat]),
        ssi_str=fmt_cont(df.loc[mask_ssi, feat]),
        p_ssi=fmt_p(test_cont(df.loc[mask_non, feat], df.loc[mask_ssi, feat])),
        missing_n=miss,
    )

# ── 3b. Binary ──
for feat in bin_cols:
    s = df[feat]
    miss = int(s.isna().sum())
    # 强制转为数值,识别 1 类
    s_num = pd.to_numeric(s, errors="coerce")
    pos_val = 1  # 假定 1 = 阳性
    n_all = (s_num.dropna()).count()
    add_row(
        feat, "binary  (n, %)",
        all_str=fmt_np((s_num == pos_val).sum(), n_all),
        dev_str=fmt_np(((s_num == pos_val) & mask_dev).sum(),
                       df.loc[mask_dev, feat].notna().sum()),
        val_str=fmt_np(((s_num == pos_val) & mask_val).sum(),
                       df.loc[mask_val, feat].notna().sum()),
        p_split=fmt_p(test_cat(df.loc[mask_dev, feat], df.loc[mask_val, feat])),
        non_str=fmt_np(((s_num == pos_val) & mask_non).sum(),
                       df.loc[mask_non, feat].notna().sum()),
        ssi_str=fmt_np(((s_num == pos_val) & mask_ssi).sum(),
                       df.loc[mask_ssi, feat].notna().sum()),
        p_ssi=fmt_p(test_cat(df.loc[mask_non, feat], df.loc[mask_ssi, feat])),
        missing_n=miss,
    )

# ── 3c. Categorical (多于 2 个 level) ──
for feat in cat_cols:
    s = df[feat]
    miss = int(s.isna().sum())
    levels = sorted(s.dropna().unique(), key=lambda x: str(x))
    # 表头行 (变量名 + 整体检验)
    p_split_overall = fmt_p(test_cat(df.loc[mask_dev, feat], df.loc[mask_val, feat]))
    p_ssi_overall   = fmt_p(test_cat(df.loc[mask_non, feat], df.loc[mask_ssi, feat]))
    add_row(
        feat, f"categorical  ({len(levels)} levels)",
        all_str="", dev_str="", val_str="", p_split=p_split_overall,
        non_str="", ssi_str="", p_ssi=p_ssi_overall,
        missing_n=miss,
    )
    # 每个 level 一个子行
    for lv in levels:
        m = s == lv
        add_row(
            f"    = {lv}", "  level",
            all_str=fmt_np(m.sum(), s.notna().sum()),
            dev_str=fmt_np((m & mask_dev).sum(), df.loc[mask_dev, feat].notna().sum()),
            val_str=fmt_np((m & mask_val).sum(), df.loc[mask_val, feat].notna().sum()),
            p_split="",
            non_str=fmt_np((m & mask_non).sum(), df.loc[mask_non, feat].notna().sum()),
            ssi_str=fmt_np((m & mask_ssi).sum(), df.loc[mask_ssi, feat].notna().sum()),
            p_ssi="",
            missing_n=None,
        )

# ══════════════════════════════════════════════════════
# 4. Output
# ══════════════════════════════════════════════════════
out_df = pd.DataFrame(rows)

# CSV (UTF-8 BOM 以便 Excel 中文显示)
out_df.to_csv(OUT_CSV, index=False, encoding="utf-8-sig")
print(f"[DONE] CSV  → {OUT_CSV}")

# XLSX with formatting
try:
    with pd.ExcelWriter(OUT_XLSX, engine="openpyxl") as writer:
        out_df.to_excel(writer, sheet_name="Table1", index=False)
        ws = writer.sheets["Table1"]
        # 列宽
        for col_cells in ws.columns:
            max_len = max(len(str(c.value)) if c.value else 0 for c in col_cells)
            ws.column_dimensions[col_cells[0].column_letter].width = min(max_len + 2, 28)
        # 表头加粗
        from openpyxl.styles import Font, PatternFill, Alignment
        bold_font = Font(bold=True)
        header_fill = PatternFill("solid", fgColor="E0E0E0")
        for c in ws[1]:
            c.font = bold_font
            c.fill = header_fill
            c.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        # p 值列突出
        p_cols_idx = [i + 1 for i, col in enumerate(out_df.columns)
                      if col.startswith("p (")]
        for ri in range(2, ws.max_row + 1):
            for ci in p_cols_idx:
                val = ws.cell(row=ri, column=ci).value
                if isinstance(val, str) and (val == "<0.001" or
                                              (val.replace(".", "").replace("-", "").isdigit()
                                               and float(val) < 0.05)):
                    ws.cell(row=ri, column=ci).font = Font(bold=True, color="B22222")
    print(f"[DONE] XLSX → {OUT_XLSX}")
except ImportError:
    print("[WARN] openpyxl 未安装。pip install openpyxl 后再跑可获得 xlsx")
except Exception as e:
    print(f"[WARN] XLSX 输出失败: {e}")


# ══════════════════════════════════════════════════════
# 5. Summary
# ══════════════════════════════════════════════════════
print("\n" + "=" * 60)
print("Summary")
print("=" * 60)
print(f"  Continuous:  {len(cont_cols)} variables")
print(f"  Binary:      {len(bin_cols)} variables")
print(f"  Categorical: {len(cat_cols)} variables")
print(f"  Total rows in table: {len(out_df)} (含 n 表头行 + 分类变量 level 行)")

# 检查 80/20 划分同质性
print("\n80/20 划分同质性检查 (DEV vs VAL, p < 0.05 视为不平衡):")
unbalanced = []
for r in rows[1:]:  # skip n header
    p_str = r.get("p (DEV vs VAL)", "")
    if isinstance(p_str, str) and p_str and p_str != "—":
        try:
            p_num = 0.0005 if p_str == "<0.001" else float(p_str)
            if p_num < 0.05:
                unbalanced.append((r["Variable"], p_str))
        except ValueError:
            pass
if unbalanced:
    print(f"  ⚠️  {len(unbalanced)} 个变量 DEV vs VAL 不平衡:")
    for v, p in unbalanced:
        print(f"     {v}  p={p}")
else:
    print("  ✓ 所有变量在 DEV vs VAL 上均无显著差异 (划分公平)")

# SSI 单变量风险因子
print("\nSSI vs non-SSI 显著相关变量 (p < 0.05):")
sig = []
for r in rows[1:]:
    p_str = r.get("p (SSI vs non-SSI)", "")
    if isinstance(p_str, str) and p_str and p_str != "—":
        try:
            p_num = 0.0005 if p_str == "<0.001" else float(p_str)
            if p_num < 0.05:
                sig.append((r["Variable"], p_str))
        except ValueError:
            pass
if sig:
    for v, p in sig:
        print(f"     {v}  p={p}")
else:
    print("  (无显著变量)")

print("\n完成。")